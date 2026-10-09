import io
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment

def build_template_workbook() -> openpyxl.Workbook:
    """
    Constructs a professionally styled Excel workbook template for candidate registration.
    Includes:
      - Sheet 1: 'Candidate Register' with pre-formatted Text cells for 16-digit numbers.
      - Sheet 2: 'Template Guidelines' with required/optional column rules and accepted headers.
    """
    wb = openpyxl.Workbook()

    # --- Sheet 1: Candidate Register ---
    ws = wb.active
    ws.title = "Candidate Register"
    ws.views.sheetView[0].showGridLines = True

    # Styling definitions matching Jerusalem College of Engineering Navy theme
    navy_header_fill = PatternFill(start_color="002F66", end_color="002F66", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10, color="1E293B")
    mono_font = Font(name="Consolas", size=10, bold=True, color="002F66")
    
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )
    
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    headers = [
        "Register Number",
        "Student Name",
        "Branch",
        "Semester",
        "Course Code",
        "Date of Birth"
    ]

    ws.row_dimensions[1].height = 28
    for col_idx, header_text in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header_text)
        cell.fill = navy_header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Add explanatory comment on Register Number header
    reg_comment = Comment(
        "REQUIRED: Exactly 16 numeric digits.\nColumn cells MUST be formatted as Text (@) to prevent Excel precision loss.",
        "COE Exam Office"
    )
    reg_comment.width = 250
    reg_comment.height = 70
    ws.cell(row=1, column=1).comment = reg_comment

    # Add explanatory comment on Date of Birth header
    dob_comment = Comment(
        "REQUIRED FOR PORTAL LOGIN: Strict DD/MM/YYYY format only (e.g. 15/08/2005).\nStudents sign in to the portal with this date of birth.",
        "COE Exam Office"
    )
    dob_comment.width = 250
    dob_comment.height = 70
    ws.cell(row=1, column=6).comment = dob_comment

    # 10 Sample candidate rows with valid 16-digit register numbers
    sample_rows = [
        ("2403310910421001", "Aarav Rajan", "B.E. Computer Science and Engineering", 5, "JCS2501", "15/08/2005"),
        ("2403310910421002", "Diya Rajan", "B.Tech. Artificial Intelligence and Data Science", 5, "JAI2502", "02/11/2005"),
        ("2403310910421003", "Karthik Rajan", "B.E. Electronics and Communication Engineering", 5, "JEC2501", "27/01/2005"),
        ("2403310910421004", "Ananya Rajan", "B.Tech. Information Technology", 5, "JIT2501", "09/03/2005"),
        ("2403310910421005", "Rahul Rajan", "B.E. Mechanical Engineering", 5, "JME2501", "21/07/2005"),
        ("2403310910421006", "Sneha Kumar", "B.E. Computer Science and Engineering", 5, "JCS2501", "30/12/2004"),
        ("2403310910421007", "Vikram Kumar", "B.Tech. Artificial Intelligence and Data Science", 5, "JAI2502", "14/06/2005"),
        ("2403310910421008", "Priya Kumar", "B.E. Electronics and Communication Engineering", 5, "JEC2501", "05/09/2005"),
        ("2403310910421009", "Siddharth Kumar", "B.Tech. Information Technology", 5, "JIT2501", "18/02/2005"),
        ("2403310910421010", "Meera Kumar", "B.E. Civil Engineering", 5, "JCE2501", "23/10/2004"),
    ]

    for row_idx, row_values in enumerate(sample_rows, 2):
        ws.row_dimensions[row_idx].height = 20
        is_even = (row_idx % 2 == 0)
        row_fill = None if is_even else zebra_fill

        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill

            if col_idx == 1:
                # REGISTER NUMBER: Explicit text format '@' and string data type
                cell.number_format = "@"
                cell.value = str(val)
                cell.font = mono_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in (2, 3):
                # Student Name & Branch: Left aligned
                cell.value = str(val)
                cell.font = data_font
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx == 4:
                # Semester: Center aligned number
                cell.value = int(val) if val is not None else ""
                cell.font = data_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 6:
                # Date of Birth (DD/MM/YYYY): Center aligned, text format
                cell.value = str(val)
                cell.number_format = "@"
                cell.font = mono_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                # Course Code: Center aligned monospace/bold
                cell.value = str(val)
                cell.font = mono_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # Pre-format extra empty rows in Column A as Text (@) so user inputs stay text
    for extra_row in range(len(sample_rows) + 2, 250):
        col_a_cell = ws.cell(row=extra_row, column=1)
        col_a_cell.number_format = "@"

    # Set column widths
    col_widths = {
        "A": 25,  # Register Number
        "B": 24,  # Student Name
        "C": 45,  # Branch
        "D": 14,  # Semester
        "E": 18,  # Course Code
        "F": 18   # Date of Birth
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    # --- Sheet 2: Guidelines & Specifications ---
    ws_guide = wb.create_sheet(title="Template Guidelines")
    ws_guide.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_guide.merge_cells("A1:E1")
    ws_guide["A1"] = "JERUSALEM COLLEGE OF ENGINEERING (AUTONOMOUS), CHENNAI"
    ws_guide["A1"].font = Font(name="Georgia", size=13, bold=True, color="002F66")
    ws_guide["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_guide.row_dimensions[1].height = 24

    ws_guide.merge_cells("A2:E2")
    ws_guide["A2"] = "Office of the Controller of Examinations — Candidate Register Template Specifications"
    ws_guide["A2"].font = Font(name="Segoe UI", size=10, bold=True, color="B8892B")
    ws_guide["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws_guide.row_dimensions[2].height = 20

    ws_guide.row_dimensions[3].height = 10  # blank

    # Table Header
    guide_headers = ["Column Header", "Requirement", "Accepted Synonyms / Aliases", "Data Type & Format", "Description / Notes"]
    guide_header_fill = PatternFill(start_color="004A99", end_color="004A99", fill_type="solid")
    guide_header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")

    ws_guide.row_dimensions[4].height = 26
    for c_idx, h_text in enumerate(guide_headers, 1):
        c = ws_guide.cell(row=4, column=c_idx, value=h_text)
        c.fill = guide_header_fill
        c.font = guide_header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    # Guidelines rows
    guide_rows = [
        (
            "Register Number",
            "REQUIRED (*)",
            "register number, reg no, reg_no, regno, registration number, register",
            "Text (@) — 16 Digits",
            "Exact 16-digit numeric string (e.g. 2403310910421001). MUST be formatted as Text in Excel to prevent 15-digit precision corruption or scientific notation (2.4033E+15)."
        ),
        (
            "Student Name",
            "Optional",
            "name, student name, student_name, candidate name",
            "Text",
            "Full legal name of the examination candidate (e.g. Aarav Rajan)."
        ),
        (
            "Branch",
            "Optional (Recommended)",
            "branch, dept, department, course, program",
            "Text",
            "Academic degree or department (e.g. B.E. Computer Science and Engineering, CSE). Highly recommended for multi-department hall interleaving."
        ),
        (
            "Semester",
            "Optional",
            "sem, semester, current semester",
            "Integer (1–8)",
            "Current student semester (e.g. 5)."
        ),
        (
            "Course Code",
            "Optional",
            "subject, subject code, subject_code, course code, course_code",
            "Text",
            "Examination subject / course code (e.g. JCS2501)."
        ),
        (
            "Date of Birth",
            "REQUIRED for portal login",
            "date of birth, dob, birth date, date_of_birth, birthdate",
            "Text — DD/MM/YYYY",
            "Student login password for the portal (e.g. 15/08/2005). Strict DD/MM/YYYY only — no other formats are accepted."
        )
    ]

    for r_idx, g_vals in enumerate(guide_rows, 5):
        ws_guide.row_dimensions[r_idx].height = 36
        for c_idx, val in enumerate(g_vals, 1):
            cell = ws_guide.cell(row=r_idx, column=c_idx, value=val)
            cell.border = thin_border
            cell.font = data_font
            if c_idx == 1:
                cell.font = Font(name="Segoe UI", size=10, bold=True, color="002F66")
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif c_idx == 2:
                # Required badge styling
                if "REQUIRED" in val:
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color="B91C1C")
                    cell.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                else:
                    cell.font = Font(name="Segoe UI", size=9, bold=False, color="166534")
                    cell.fill = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 3:
                cell.font = Font(name="Consolas", size=9, color="475569")
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            elif c_idx == 4:
                cell.font = Font(name="Segoe UI", size=9, bold=True, color="1E293B")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # Advisory Section
    adv_start = len(guide_rows) + 6
    ws_guide.merge_cells(f"A{adv_start}:E{adv_start}")
    ws_guide[f"A{adv_start}"] = "CRITICAL EXCEL FORMATTING ADVISORY FOR REGISTER NUMBERS"
    ws_guide[f"A{adv_start}"].font = Font(name="Segoe UI", size=11, bold=True, color="92400E")
    ws_guide[f"A{adv_start}"].fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    ws_guide[f"A{adv_start}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_guide.row_dimensions[adv_start].height = 26

    advisory_notes = [
        "1. Excel 15-Digit IEEE-754 Limit: Microsoft Excel only supports 15 digits for numbers. A 16-digit number saved as General or Number will truncate the 16th digit to '0' or convert into scientific notation (e.g., 2.40331E+15).",
        "2. How to ensure proper Text format in Excel: Select Column A -> Right click -> Format Cells -> choose 'Text' -> click OK before entering or pasting register numbers.",
        "3. Alternative Quick Fix: Precede the 16 digits with a single apostrophe (e.g., '2403310910421001). The system automatically removes the leading apostrophe during import.",
        "4. Supported Upload Formats: Both Microsoft Excel (.xlsx, .xls) and Comma-Separated Values (.csv) are fully supported by the validation engine."
    ]

    for note_idx, note_text in enumerate(advisory_notes, adv_start + 1):
        ws_guide.merge_cells(f"A{note_idx}:E{note_idx}")
        cell = ws_guide[f"A{note_idx}"]
        cell.value = note_text
        cell.font = Font(name="Segoe UI", size=9, italic=False, color="78350F")
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
        ws_guide.row_dimensions[note_idx].height = 24

    # Column widths for Guidelines sheet
    ws_guide.column_dimensions["A"].width = 20
    ws_guide.column_dimensions["B"].width = 24
    ws_guide.column_dimensions["C"].width = 38
    ws_guide.column_dimensions["D"].width = 24
    ws_guide.column_dimensions["E"].width = 50

    return wb

def generate_candidate_template_xlsx() -> io.BytesIO:
    """
    Generates the template workbook and returns an in-memory BytesIO buffer.
    """
    wb = build_template_workbook()
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

def save_template_files(project_root: Path = None):
    """
    Saves the generated candidate_register_template.xlsx to:
      1. samples/candidate_register_template.xlsx
      2. frontend/public/candidate_register_template.xlsx
    """
    if project_root is None:
        project_root = Path(__file__).resolve().parent.parent.parent

    wb = build_template_workbook()

    samples_path = project_root / "samples" / "candidate_register_template.xlsx"
    samples_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(samples_path))

    public_path = project_root / "frontend" / "public" / "candidate_register_template.xlsx"
    public_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(public_path))

if __name__ == "__main__":
    save_template_files()
    print("Candidate register template files updated successfully.")
